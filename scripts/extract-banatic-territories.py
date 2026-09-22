#!/usr/bin/env python3
import csv
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

def norm(value):
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()

def canonical_type(value):
    raw = norm(value)
    if raw == "cc" or "communaute de communes" in raw:
        return "CC"
    if raw == "ca" or "communaute d agglomeration" in raw:
        return "CA"
    return None

ALIASES = {
    "siren": ["n siren", "numero siren", "siren", "code officiel groupement"],
    "name": ["nom du groupement", "groupement", "denomination"],
    "type": ["nature juridique", "code nature juridique", "nature"],
    "department": ["departement", "dep"],
    "arrondissement": ["arrondissement siege", "arrondissement"],
    "seat": ["commune siege", "commune du siege"],
    "population": ["population totale", "population totale regroupee", "population"],
    "members": ["nombre de membres", "nb membres"],
    "president_title": ["civilite president", "civilite du president"],
    "president_last": ["nom president", "nom du president"],
    "president_first": ["prenom president", "prenom du president"],
    "address1": ["numero et libelle de la voie du siege", "adresse du siege", "adresse siege"],
    "address2": ["complement d adresse du siege", "complement adresse siege"],
    "postal": ["code postal du siege", "code postal siege"],
    "city": ["ville du siege", "ville siege"],
    "phone": ["telephone du siege", "telephone siege", "telephone"],
    "email": ["courriel du siege", "adresse electronique", "courriel", "email"],
    "website": ["site internet du siege", "site internet", "site web"]
}

def get(row, headers, key):
    for alias in ALIASES[key]:
        if alias in headers:
            value = row[headers[alias]]
            if value not in (None, ""):
                return value
    return None

def integer(value):
    if value in (None, ""):
        return None
    digits = re.sub(r"[^0-9-]", "", str(value))
    try:
        return int(digits)
    except Exception:
        return None

def siren(value):
    digits = re.sub(r"\D", "", "" if value is None else str(value))
    return digits.zfill(9) if digits else None

def make_item(row, headers):
    territory_type = canonical_type(get(row, headers, "type"))
    if territory_type not in {"CC", "CA"}:
        return None
    code = siren(get(row, headers, "siren"))
    name = get(row, headers, "name")
    if not code or len(code) != 9 or not name:
        return None

    dep = get(row, headers, "department")
    dep_match = re.search(r"\b(\d{2,3}|2A|2B)\b", "" if dep is None else str(dep), re.I)
    department_code = dep_match.group(1).upper() if dep_match else (str(dep).strip() if dep else None)

    return {
        "siren": code,
        "name": str(name).strip(),
        "territoryType": territory_type,
        "departmentCode": department_code,
        "arrondissement": str(get(row, headers, "arrondissement") or "").strip() or None,
        "seatCommune": str(get(row, headers, "seat") or "").strip() or None,
        "populationTotal": integer(get(row, headers, "population")),
        "memberCount": integer(get(row, headers, "members")),
        "presidentTitle": str(get(row, headers, "president_title") or "").strip() or None,
        "presidentLastName": str(get(row, headers, "president_last") or "").strip() or None,
        "presidentFirstName": str(get(row, headers, "president_first") or "").strip() or None,
        "addressLine1": str(get(row, headers, "address1") or "").strip() or None,
        "addressLine2": str(get(row, headers, "address2") or "").strip() or None,
        "postalCode": str(get(row, headers, "postal") or "").strip() or None,
        "city": str(get(row, headers, "city") or "").strip() or None,
        "phone": str(get(row, headers, "phone") or "").strip() or None,
        "email": str(get(row, headers, "email") or "").strip() or None,
        "website": str(get(row, headers, "website") or "").strip() or None,
        "banaticUrl": f"https://www.banatic.interieur.gouv.fr/intercommunalite/{code}",
        "sourcePayload": {"source": "BANATIC/DGCL"}
    }

def merge(existing, incoming):
    if not existing:
        return incoming
    for key, value in incoming.items():
        if existing.get(key) in (None, "", 0) and value not in (None, "", 0):
            existing[key] = value
    return existing

def parse_csv(path):
    items = {}
    raw = Path(path).read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        raise RuntimeError("Unable to decode BANATIC CSV")

    sample = text[:10000]
    dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    reader = csv.reader(text.splitlines(), dialect)
    rows = list(reader)
    if not rows:
        return items
    header_row = rows[0]
    headers = {norm(v): idx for idx, v in enumerate(header_row)}
    for values in rows[1:]:
        row = list(values) + [None] * max(0, len(header_row) - len(values))
        item = make_item(row, headers)
        if item:
            items[item["siren"]] = merge(items.get(item["siren"]), item)
    return items

def parse_xlsx(path):
    from openpyxl import load_workbook
    items = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        header = None
        headers = None
        for row in ws.iter_rows(values_only=True):
            values = list(row)
            normalized = [norm(v) for v in values]
            if header is None:
                has_siren = any(v in ALIASES["siren"] for v in normalized)
                has_name = any(v in ALIASES["name"] for v in normalized)
                has_type = any(v in ALIASES["type"] for v in normalized)
                if has_siren and has_name and has_type:
                    header = values
                    headers = {norm(v): idx for idx, v in enumerate(values) if v not in (None, "")}
                continue
            item = make_item(values, headers)
            if item:
                items[item["siren"]] = merge(items.get(item["siren"]), item)
    return items

def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: extract-banatic-territories.py INPUT OUTPUT")
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    if source.suffix.lower() == ".csv":
        items = parse_csv(source)
    else:
        items = parse_xlsx(source)

    values = sorted(items.values(), key=lambda x: (x["territoryType"], x["name"]))
    counts = {
        "CC": sum(1 for x in values if x["territoryType"] == "CC"),
        "CA": sum(1 for x in values if x["territoryType"] == "CA")
    }
    if not values:
        raise RuntimeError("BANATIC extraction returned zero CC/CA territories")

    payload = {
        "sourceUpdatedAt": datetime.now(timezone.utc).isoformat(),
        "items": values,
        "counts": counts
    }
    output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"total": len(values), **counts}, ensure_ascii=False))

if __name__ == "__main__":
    main()
